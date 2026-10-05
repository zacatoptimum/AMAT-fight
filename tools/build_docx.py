from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT_NAME = "Microsoft JhengHei"
MONO_FONT = "Consolas"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
PALE_GRAY = "F5F7F9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_GRAY, size="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def remove_paragraph_borders(element) -> None:
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_run_font(run, name=FONT_NAME, size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline_runs(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos:match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), bold=True)
        else:
            set_run_font(paragraph.add_run(token[1:-1]), name=MONO_FONT, size=9.5)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, size=8, color="666666")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    style_specs = {
        "Title": (25, 16, 8),
        "Heading 1": (16, 14, 6),
        "Heading 2": (13, 12, 4),
        "Heading 3": (11.5, 10, 3),
    }
    for name, (size, before, after) in style_specs.items():
        style = doc.styles[name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        remove_paragraph_borders(style._element)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        stripped = line.strip().strip("|")
        rows.append([cell.strip() for cell in stripped.split("|")])
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = True
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        # Repeat and keep only the header row together. Body rows may split:
        # narrow, text-heavy tables can otherwise trigger an endless Word
        # pagination loop during PDF export even when character counts look low.
        if i == 0:
            row_props = table.rows[i]._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            row_props.append(cant_split)
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            text = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            add_inline_runs(p, text)
            for run in p.runs:
                set_run_font(run, size=8.7, color="FFFFFF" if i == 0 else "000000", bold=i == 0)
            if i == 0:
                set_cell_shading(cell, NAVY)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_docx(source: Path, output: Path) -> None:
    text = source.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    doc = Document()
    style_document(doc)

    i = 0
    in_code = False
    code_lines: list[str] = []
    title_written = False
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, name=MONO_FONT, size=8.5, color="333333")
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.strip() == "<!-- PAGEBREAK -->":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].lstrip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, parse_table(table_lines))
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            content = heading.group(2).strip()
            if level == 1 and not title_written:
                p = doc.add_paragraph(style="Title")
                p.paragraph_format.space_after = Pt(10)
                remove_paragraph_borders(p._p)
                add_inline_runs(p, content)
                for run in p.runs:
                    set_run_font(run, size=25, color="000000", bold=False)
                title_written = True
            else:
                p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
                add_inline_runs(p, content)
                for run in p.runs:
                    set_run_font(run, color="000000", bold=True)
            i += 1
            continue
        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline_runs(p, bullet.group(1))
            i += 1
            continue
        number = re.match(r"^\s*(\d+)[.)]\s+(.+)$", line)
        if number:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            p.paragraph_format.space_after = Pt(3)
            prefix = p.add_run(f"{number.group(1)}.  ")
            set_run_font(prefix, bold=False)
            add_inline_runs(p, number.group(2))
            i += 1
            continue

        paragraph_lines = [line.strip()]
        i += 1
        while i < len(lines):
            next_line = lines[i].rstrip()
            if (
                not next_line.strip()
                or next_line.startswith("#")
                or next_line.lstrip().startswith("|")
                or next_line.strip().startswith("```")
                or next_line.strip() == "<!-- PAGEBREAK -->"
                or re.match(r"^\s*[-*]\s+", next_line)
                or re.match(r"^\s*\d+[.)]\s+", next_line)
            ):
                break
            paragraph_lines.append(next_line.strip())
            i += 1
        p = doc.add_paragraph()
        add_inline_runs(p, " ".join(paragraph_lines))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = lines[0].lstrip("# ").strip() if lines else source.stem
    doc.core_properties.subject = "AMAT Pre-delivery Repository v0.1"
    doc.core_properties.author = "Optimum"
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pairs", nargs="+", help="source.md=output.docx")
    args = parser.parse_args()
    for pair in args.pairs:
        source_raw, output_raw = pair.split("=", 1)
        build_docx(Path(source_raw), Path(output_raw))
        print(f"BUILT {output_raw}")


if __name__ == "__main__":
    main()
